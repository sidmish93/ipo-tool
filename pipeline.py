"""
IPO shareholding Excel pipeline.

Given a date range, this:
  1. Lists mainboard IPOs by listing date from Chittorgarh's timetable.
  2. Looks up each on BSE (scrip code, tickers, ISIN, market cap).
  3. Keeps companies with live Mcap Full above a chosen threshold (default 3000 cr).
  4. Finds the latest shareholding quarter + statement page links.
  5. Extracts detailed promoter and public shareholding.
  6. Writes a two-sheet Excel workbook.

Exchange pages run inside a real Chrome browser (via Playwright) because
BSE/NSE block plain HTTP and bundled Chromium.
"""

import os
import re
import json
import time
import datetime
import urllib.parse
import urllib.request
import urllib.error

from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from deals import fetch_and_adjust, quarter_end as _quarter_end_date

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/142.0.0.0 Safari/537.36")

# BSE's WAF returns 403 to Playwright's bundled Chromium. Use a real
# Chrome/Edge install and verify BSE actually loads before continuing.
_BROWSER_ARGS = [
    "--disable-blink-features=AutomationControlled",
    "--no-sandbox",
    "--disable-setuid-sandbox",
    "--disable-dev-shm-usage",
    "--disable-gpu",
    "--disable-extensions",
    "--disable-background-networking",
    "--renderer-process-limit=3",
    "--mute-audio",
    "--no-first-run",
    "--no-default-browser-check",
]


def _low_mem():
    """True on Render / 512 MB boxes — keep Chrome off except one company."""
    flag = (os.environ.get("LOW_MEM") or "").strip().lower()
    if flag in ("1", "true", "yes"):
        return True
    if flag in ("0", "false", "no"):
        return False
    if os.environ.get("RENDER"):
        return True
    try:
        with open("/proc/meminfo", encoding="ascii") as fh:
            for line in fh:
                if line.startswith("MemTotal:"):
                    kb = int(line.split()[1])
                    return kb < 768 * 1024
    except (OSError, ValueError, IndexError):
        pass
    return False


def _chrome_args():
    args = list(_BROWSER_ARGS)
    if _low_mem():
        args.extend(["--renderer-process-limit=1"])
    return args


def _short_error(exc):
    text = str(exc).replace("\r", "\n").strip()
    first = next((ln.strip() for ln in text.splitlines() if ln.strip()),
                 type(exc).__name__)
    if len(first) > 240:
        first = first[:237] + "..."
    return first
_STEALTH_JS = "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"


def _playwright_chrome_bins():
    """Chrome-for-Testing binaries under Playwright's cache (Render/Docker)."""
    roots = [
        os.environ.get("PLAYWRIGHT_BROWSERS_PATH"),
        os.path.expanduser("~/.cache/ms-playwright"),
        "/ms-playwright",
        "/opt/render/project/.cache/ms-playwright",
        "/opt/render/.cache/ms-playwright",
    ]
    found = []
    for root in roots:
        if not root or not os.path.isdir(root):
            continue
        try:
            for dirpath, dirnames, files in os.walk(root):
                low = dirpath.replace("\\", "/").lower()
                if "chromium" in low:
                    dirnames[:] = []
                    continue
                if "chrome-" not in low and "/chrome/" not in low:
                    continue
                for name in files:
                    if name in ("chrome", "chrome.exe"):
                        found.append(os.path.join(dirpath, name))
        except OSError:
            continue
    return found


def _installed_browsers():
    """Yield (label, executable_path) for Chrome/Edge on this machine."""
    home = os.path.expanduser("~")
    candidates = [
        ("Chrome", os.environ.get("CHROME_PATH") or os.environ.get("CHROME_BIN")),
        ("Chrome", os.path.join(os.environ.get("PROGRAMFILES", r"C:\Program Files"),
                                "Google", "Chrome", "Application", "chrome.exe")),
        ("Chrome", os.path.join(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)"),
                                "Google", "Chrome", "Application", "chrome.exe")),
        ("Chrome", os.path.join(os.environ.get("LOCALAPPDATA", ""),
                                "Google", "Chrome", "Application", "chrome.exe")),
        ("Edge", os.path.join(os.environ.get("PROGRAMFILES", r"C:\Program Files"),
                              "Microsoft", "Edge", "Application", "msedge.exe")),
        ("Edge", os.path.join(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)"),
                              "Microsoft", "Edge", "Application", "msedge.exe")),
        ("Chrome", "/usr/bin/google-chrome"),
        ("Chrome", "/usr/bin/google-chrome-stable"),
        ("Chrome", "/opt/google/chrome/chrome"),
        ("Chrome", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        ("Edge", "/usr/bin/microsoft-edge"),
        ("Edge", os.path.join(home, "AppData", "Local", "Microsoft", "Edge",
                              "Application", "msedge.exe")),
    ]
    for path in _playwright_chrome_bins():
        candidates.append(("Chrome", path))
    seen = set()
    for label, path in candidates:
        if path and os.path.isfile(path) and path not in seen:
            seen.add(path)
            yield label, path


def _can_open_window():
    if os.name == "nt":
        return True
    return bool(os.environ.get("DISPLAY"))


def _browser_specs():
    """Launch recipes. Real Chrome/Edge only — bundled Chromium is blocked by BSE."""
    common = {
        "args": _chrome_args(),
        "ignore_default_args": ["--enable-automation"],
    }
    specs = []
    for label, path in _installed_browsers():
        specs.append({
            "label": f"{label} headless",
            "launch": {"executable_path": path, "headless": True, **common},
        })
    specs.append({"label": "Playwright channel=chrome",
                  "launch": {"channel": "chrome", "headless": True, **common}})
    specs.append({"label": "Playwright channel=msedge",
                  "launch": {"channel": "msedge", "headless": True, **common}})
    if _can_open_window():
        for label, path in _installed_browsers():
            specs.append({
                "label": f"{label} window",
                "launch": {"executable_path": path, "headless": False, **common},
            })
    return specs


def _new_context(browser):
    ctx = browser.new_context(
        user_agent=UA,
        viewport={"width": 1440, "height": 900},
        locale="en-IN",
        extra_http_headers={"Accept-Language": "en-IN,en;q=0.9"},
    )
    ctx.add_init_script(_STEALTH_JS)
    ctx.set_default_timeout(60000)
    return ctx


def _bse_blocked(page) -> bool:
    try:
        title = (page.title() or "").lower()
    except Exception:
        title = ""
    return "access denied" in title


def _bse_reachable(ctx) -> bool:
    page = ctx.new_page()
    try:
        resp = page.goto("https://www.bseindia.com/",
                         wait_until="domcontentloaded", timeout=45000)
        page.wait_for_timeout(1200)
        status = resp.status if resp else 0
        return status == 200 and not _bse_blocked(page)
    except Exception:
        return False
    finally:
        try:
            page.close()
        except Exception:
            pass


def _open_session(playwright, probe=True):
    """Launch a browser that BSE will actually talk to.

    Returns (browser, context, label). `probe=False` skips the extra BSE
    tab (needed on 512 MB so we only ever have one Chrome page).
    """
    errors = []
    for spec in _browser_specs():
        try:
            browser = playwright.chromium.launch(**spec["launch"])
        except Exception as e:
            errors.append(f"{spec['label']}: could not start ({e})")
            continue
        ctx = _new_context(browser)
        if not probe or _bse_reachable(ctx):
            return browser, ctx, spec["label"]
        try:
            ctx.close()
        except Exception:
            pass
        try:
            browser.close()
        except Exception:
            pass
        errors.append(f"{spec['label']}: BSE returned Access Denied")
    detail = " | ".join(errors[:6]) if errors else "no Chrome/Edge found"
    raise RuntimeError(
        "BSE blocked the browser. On a PC, install Google Chrome (not only "
        "Playwright) and close extra Chrome windows. On Render, the service "
        "must use the Docker environment so the image can install real Chrome "
        "(native Python + Playwright Chromium is blocked). "
        f"Tried: {detail}"
    )

SEBI_LIST = ("https://www.sebi.gov.in/sebiweb/home/HomeAction.do"
             "?doListing=yes&sid=3&ssid=15&smid=12")
CHITTOR_HOME = "https://www.chittorgarh.com/"
# Report 118 (timetable) is the full year list. Report 25 (listing-date)
# paywalls older years at 5 rows.
CHITTOR_LIST_API = (
    "https://webnodejs.chittorgarh.com/cloud/report/data-read/"
    "118/1/9/{year}/0/0/mainboard/0?search=&v=13-18"
)
MCAP_THRESHOLD = 3000.0


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def _norm(name):
    """Normalise a company name for fuzzy matching."""
    if not name:
        return ""
    s = name.lower()
    s = re.sub(r"\band\b", " ", s)
    s = re.sub(r"[^a-z0-9]", "", s)
    for suf in ("privatelimited", "limited", "ltd", "pvt"):
        if s.endswith(suf):
            s = s[: -len(suf)]
    return s


_BSE_SKIP_WORDS = {
    "&", "and", "the", "of", "ltd", "ltd.", "limited", "pvt", "private",
}


def _bse_search_queries(name):
    """Query strings BSE's search actually answers.

    Their index strips '&' (so 'Larsen & Toubro' is stored as
    'Larsen  Toubro Ltd') and a full-name search returns nothing. First
    significant word ('Larsen') and the ticker-like form still hit.
    """
    stripped = re.sub(r"\s+(limited|ltd|pvt|private)\.?$", "", name or "",
                      flags=re.IGNORECASE)

    def significant(text):
        return [w for w in re.split(r"[\s,/]+", (text or "").strip())
                if w and w.lower() not in _BSE_SKIP_WORDS]

    raw = []
    for base in (name, stripped):
        if not base:
            continue
        raw.append(base)
        raw.append(re.sub(r"\s*&\s*", " ", base))
        raw.append(re.sub(r"\s+and\s+", " ", base, flags=re.IGNORECASE))
        words = significant(base)
        if len(words) >= 3:
            raw.append(" ".join(words[:3]))
        if len(words) >= 2:
            raw.append(" ".join(words[:2]))
        if words and len(words[0]) >= 3:
            raw.append(words[0])
        if re.search(r"\bAMC\b", base or "", re.I):
            raw.append(re.sub(r"\bAMC\b", "Asset Management", base, flags=re.I))
        if re.fullmatch(r"LIC", (base or "").strip(), re.I):
            raw.append("Life Insurance Corporation")

    seen, out = set(), []
    for q in raw:
        q = re.sub(r"\s+", " ", q or "").strip(" &")
        key = q.lower()
        if q and key not in seen:
            seen.add(key)
            out.append(q)
    return out


def _to_float(txt):
    if txt is None:
        return None
    s = str(txt).replace(",", "").strip()
    if s in ("", "-", "NA", "N.A."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


_MONTHS = {m[:3].lower(): m for m in
           ["January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December"]}


def _display_quarter(qname):
    """Normalise a BSE quarter label to 'Month YYYY'."""
    if not qname:
        return ""
    qname = qname.strip()
    m = re.match(r"^\d{1,2}[-/ ]([A-Za-z]{3,})[-/ ](\d{2,4})$", qname)
    if m:
        mon = _MONTHS.get(m.group(1)[:3].lower(), m.group(1).title())
        yr = m.group(2)
        if len(yr) == 2:
            yr = "20" + yr
        return f"{mon} {yr}"
    return qname


def split_terms(text):
    """Company names as typed: several at a time, separated by semicolons.

    Same rule as sidmish93/blocks_tracker: 'Delhivery; Lodha; Vedanta'.
    """
    seen, terms = set(), []
    for part in (text or "").split(";"):
        term = part.strip()
        if term and term.lower() not in seen:
            seen.add(term.lower())
            terms.append(term)
    return terms


def _holding_cr(pct, mcap):
    """Rupee size of a holding in crores: shareholding % × live mcap."""
    if pct is None or mcap is None:
        return None
    try:
        return round(float(pct) * float(mcap) / 100.0, 2)
    except (TypeError, ValueError):
        return None


def _with_holdings(rows, mcap):
    out = []
    for row in rows or []:
        if isinstance(row, dict):
            row["holding_cr"] = _holding_cr(row.get("pct"), mcap)
            out.append(row)
            continue
        if not row or len(row) < 3:
            continue
        nm, extra, pct = row[0], row[1], row[2]
        shares = row[3] if len(row) > 4 else None
        out.append({
            "name": nm,
            "category": extra,
            "pct": pct,
            "shares": shares,
            "holding_cr": _holding_cr(pct, mcap),
            "bought": 0,
            "sold": 0,
            "adj_shares": shares,
            "adj_pct": pct,
            "adj_holding_cr": _holding_cr(pct, mcap),
        })
    return out


def _iso(d):
    if d is None:
        return ""
    if hasattr(d, "isoformat"):
        return d.isoformat()
    return str(d)


def _chittor_listing_date(row):
    for key in ("~IL_IPO_Listing_date", "~IPO_Listing_date"):
        raw = str(row.get(key) or "").strip()
        if raw:
            try:
                return datetime.datetime.fromisoformat(raw[:10]).date()
            except ValueError:
                pass
    text = str(row.get("Listing Date") or "").strip()
    for fmt in ("%d-%b-%Y", "%d-%B-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _clean_company(title):
    """'Turtlemint Fintech Solutions Limited - Prospectus' -> company name.

    Some SME rows contain two lines (Prospectus + Abridged Prospectus); keep the
    first line only, then strip the trailing '- Prospectus' marker.
    """
    first_line = title.splitlines()[0] if title else ""
    first_line = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", first_line)
    t = re.sub(r"\s*[-–—]\s*(abridged\s+)?prospectus.*$", "", first_line,
               flags=re.IGNORECASE).strip()
    return t


# --------------------------------------------------------------------------- #
# Browser-evaluated JS snippets
# --------------------------------------------------------------------------- #
_SEBI_ROWS_JS = r"""
() => {
  const rows = [];
  document.querySelectorAll('table tr').forEach(tr => {
    const td = tr.querySelectorAll('td');
    if (td.length >= 2) {
      const d = (td[0].innerText || '').trim();
      const t = (td[1].innerText || '').trim();
      if (t && /\d{4}/.test(d)) rows.push([d, t]);
    }
  });
  return rows;
}
"""

# --------------------------------------------------------------------------- #
# Pipeline
# --------------------------------------------------------------------------- #
class Cancelled(Exception):
    """Raised when the user stops a date-range run."""


class Pipeline:
    def __init__(self, progress_cb=None, cancel_cb=None):
        self.cb = progress_cb or (lambda *a, **k: None)
        self.cancel_cb = cancel_cb or (lambda: False)

    def log(self, msg, current=None, total=None, stage=None, **extra):
        self.cb(msg, current=current, total=total, stage=stage, **extra)

    def _check_cancel(self):
        if self.cancel_cb():
            raise Cancelled("Cancelled.")

    # ----- IPO name list (Chittorgarh listing date) ---------------------- #
    def _listed_companies(self, page, dfrom, dto):
        """Mainboard IPOs whose listing date falls in [dfrom, dto].

        Year tabs are by issue year, so a Dec open / Jan listing sits on
        the previous tab. Fetch year-1 through the end year, then keep
        rows whose listing date is in range. Skip names with no listing
        date yet (still in the issue window).
        """
        out, seen = [], set()
        if page is not None:
            try:
                page.goto(CHITTOR_HOME, wait_until="domcontentloaded",
                          timeout=45000)
                page.wait_for_timeout(600)
            except Exception:
                pass
        for year in range(dfrom.year - 1, dto.year + 1):
            self._check_cancel()
            rows = self._chittor_year(page, year)
            kept = 0
            for name, dt in rows:
                if not (dfrom <= dt <= dto):
                    continue
                key = _norm(name)
                if not key or key in seen:
                    continue
                seen.add(key)
                out.append((name, dt))
                kept += 1
            self.log(f"{year}: {len(rows)} timetable rows, {kept} listed "
                     f"in range.", stage="list")
        out.sort(key=lambda x: x[1])
        return out

    def _rows_from_chittor_payload(self, raw):
        rows = []
        for item in raw or []:
            name = re.sub(r"<[^>]+>", "", str(item.get("Company") or ""))
            name = re.sub(r"\s+", " ", name).strip(" .")
            dt = _chittor_listing_date(item)
            if name and dt:
                rows.append((name, dt))
        return rows

    def _chittor_year_http(self, year):
        url = CHITTOR_LIST_API.format(year=year)
        req = urllib.request.Request(url, headers={
            "User-Agent": UA,
            "Accept": "application/json, text/plain, */*",
            "Referer": CHITTOR_HOME,
            "Origin": "https://www.chittorgarh.com",
        })
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                payload = json.loads(resp.read().decode("utf-8", "replace"))
        except (urllib.error.URLError, TimeoutError, ValueError, OSError):
            return []
        if not isinstance(payload, dict):
            return []
        return self._rows_from_chittor_payload(payload.get("reportTableData"))

    def _chittor_year(self, page, year):
        """[(name, listing_date), ...] from the mainboard timetable."""
        if page is None:
            return self._chittor_year_http(year)
        url = CHITTOR_LIST_API.format(year=year)
        res = None
        try:
            res = page.evaluate(self._FETCH_JS, url)
        except Exception:
            try:
                page.goto(CHITTOR_HOME, wait_until="domcontentloaded",
                          timeout=45000)
                page.wait_for_timeout(400)
                res = page.evaluate(self._FETCH_JS, url)
            except Exception:
                res = None
        if res and res.get("ok") and isinstance(res.get("json"), dict):
            return self._rows_from_chittor_payload(
                res["json"].get("reportTableData"))
        return []

    def _sebi_companies(self, page, dfrom, dto):
        """Return [(name, date)] of prospectuses filed within [dfrom, dto]."""
        page.goto(SEBI_LIST, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(1500)
        f_str = dfrom.strftime("%d-%m-%Y")
        t_str = dto.strftime("%d-%m-%Y")
        page.evaluate(
            "(v)=>{document.getElementById('fromDate').value=v.f;"
            "document.getElementById('toDate').value=v.t;}",
            {"f": f_str, "t": t_str},
        )

        def trigger(js):
            try:
                with page.expect_response(
                    lambda r: "HomeAction.do" in r.url
                    and r.request.method == "POST",
                    timeout=25000,
                ):
                    page.evaluate(js)
            except Exception:
                page.evaluate(js)
            page.wait_for_timeout(1200)

        all_rows = []
        prev_first = None
        for page_idx in range(0, 25):
            if page_idx == 0:
                trigger("()=>searchFormNewsList('s','-1')")
            else:
                trigger(f"()=>searchFormNewsList('n','{page_idx}')")
            rows = page.evaluate(_SEBI_ROWS_JS)
            if not rows:
                break
            first = tuple(rows[0])
            if page_idx > 0 and first == prev_first:
                break  # pager clamped -> no more pages
            all_rows.extend(rows)
            prev_first = first
            self.log(f"SEBI page {page_idx + 1}: {len(rows)} rows",
                     stage="sebi")
            has_next = page.evaluate(
                r"""() => {
                  let n=false;
                  document.querySelectorAll('a').forEach(a=>{
                    const t=(a.innerText||'').trim().toLowerCase();
                    const h=a.getAttribute('href')||'';
                    if((t==='next'||t==='last')&&/searchFormNewsList/.test(h)) n=true;
                  });
                  return n;
                }"""
            )
            if not has_next:
                break

        # parse + filter by date + dedupe
        out = []
        seen = set()
        for d, t in all_rows:
            try:
                dt = datetime.datetime.strptime(d.strip(), "%b %d, %Y").date()
            except ValueError:
                continue
            if not (dfrom <= dt <= dto):
                continue
            name = _clean_company(t)
            key = _norm(name)
            if not key or key in seen:
                continue
            seen.add(key)
            out.append((name, dt))
        return out

    # ----- BSE ------------------------------------------------------------ #
    _FETCH_JS = """async (u) => {
        try {
          const r = await fetch(u, {headers:{'Accept':'application/json'}});
          const t = await r.text();
          try { return {ok:true, json:JSON.parse(t)}; }
          catch(e){ return {ok:false, text:t.slice(0,200)}; }
        } catch(e){ return {ok:false, text:''+e}; }
      }"""

    def _bse_fetch(self, bpage, url, retries=3):
        for attempt in range(retries):
            try:
                return bpage.evaluate(self._FETCH_JS, url)
            except Exception:
                try:
                    bpage.wait_for_timeout(700)
                except Exception:
                    pass
        return {"ok": False, "text": "evaluate failed"}

    @staticmethod
    def _seg_rank(type_str):
        """Rank BSE search segments: Equity T+1 is the one to use.

        A company can appear under 'Equity T+1', 'Derivatives' and 'Equity T+0'.
        The T+0 row has a *different* scrip code (e.g. 143529 / DELHIVERY#) whose
        shareholding pattern is blank, so it must be avoided in favour of T+1.
        """
        t = (type_str or "").lower()
        if "equity t+1" in t:
            return 0
        if "equity t+0" in t:
            return 3
        if "deriv" in t:
            return 4
        if "equity" in t:
            return 1
        return 2

    @staticmethod
    def _name_rank(scripname, target, shortname="", query_name=""):
        n = _norm(scripname)
        t = _norm(shortname)
        if n and n == target:
            return 0
        if t and t == target:
            return 0

        q_words = [w.lower() for w in re.split(r"[\s,/&]+", query_name or "")
                   if w and w.lower() not in _BSE_SKIP_WORDS and len(w) >= 3]
        s_words = [w.lower() for w in re.split(r"[\s,/&]+",
                   f"{scripname or ''} {shortname or ''}")
                   if w and w.lower() not in _BSE_SKIP_WORDS]
        _TOKEN_EQ = {
            "amc": {"amc", "asset", "management"},
        }
        _SHORT_EXPAND = {
            "lic": ["life", "insurance", "corporation"],
        }

        def _hit(qw):
            aliases = _TOKEN_EQ.get(qw, {qw})
            return any(
                aw == sw or (len(aw) >= 4 and (aw in sw or sw in aw))
                for aw in aliases for sw in s_words
            )

        # "LIC" is Life Insurance Corporation, not LIC Housing / LIC MF.
        if len(q_words) == 1 and q_words[0] in _SHORT_EXPAND:
            q_words = _SHORT_EXPAND[q_words[0]]
        elif len(q_words) == 1 and len(q_words[0]) < 5:
            return 2

        missing = [w for w in q_words if not _hit(w)] if q_words else []
        if missing:
            return 2

        def _contained(a, b):
            if not a or not b:
                return False
            shorter, longer = (a, b) if len(a) <= len(b) else (b, a)
            # reject "ICICI" ⊂ "ICICIPRUDENTIALAMC" style false hits
            return shorter in longer and len(shorter) >= 8

        if _contained(n, target) or _contained(t, target):
            return 1
        if q_words and s_words and all(_hit(w) for w in q_words):
            return 1
        return 2

    def _bse_lookup(self, bpage, name):
        """Return dict with scripcode/ticker/isin/bse_name or None.

        A company can show up in several segments. Searching the *full* name
        ("X Limited") often returns ONLY the Equity T+0 phantom (scrip like
        143529 / 'X#') whose shareholding pattern is blank, while the real
        Equity T+1 listing ('X Ltd', 543529) only appears for a shorter query.
        So we collect hits across all query variants and pick the best name
        match in the Equity T+1 segment (see _seg_rank / _name_rank).
        """
        queries = _bse_search_queries(name)

        target = _norm(name)
        seen, cands = set(), []

        def rank(r):
            return (self._name_rank(r.get("scripName", ""), target,
                                    r.get("shortName", ""), name),
                    self._seg_rank(r.get("Type", "")))

        for q in queries:
            url = ("https://api.bseindia.com/BseIndiaAPI/api/"
                   "GetQuoteAllSearchDatabeta/w?searchString="
                   + urllib.parse.quote(q))
            res = self._bse_fetch(bpage, url)
            if res.get("ok") and isinstance(res["json"], list):
                for r in res["json"]:
                    key = (str(r.get("strSricpCode", "")), r.get("Type", ""))
                    if key not in seen:
                        seen.add(key)
                        cands.append(r)
            if cands:
                best = min(cands, key=rank)
                nr, sr = rank(best)
                # stop only once we have a solid name match in a real equity
                # segment (T+1 or plain equity) — never on a T+0/derivative row
                if nr <= 1 and sr <= 1:
                    return self._pack_lookup(best)

        if cands:
            best = min(cands, key=rank)
            nr, sr = rank(best)
            # Equity only — never fall back to a debt / ETF / T+0 row
            if nr <= 1 and sr <= 1:
                return self._pack_lookup(best)
        return None

    @staticmethod
    def _pack_lookup(r):
        return {
            "scripcode": str(r.get("strSricpCode", "")).strip(),
            "ticker": (r.get("shortName") or "").strip(),
            "isin": (r.get("Isin") or "").strip(),
            "bse_name": (r.get("scripName") or "").strip(),
        }

    def _bse_mcap(self, bpage, scripcode):
        url = ("https://api.bseindia.com/BseIndiaAPI/api/StockTrading/w"
               "?flag=&quotetype=EQ&scripcode=" + scripcode)
        res = self._bse_fetch(bpage, url)
        if res.get("ok"):
            return _to_float(res["json"].get("MktCapFull"))
        return None

    def _bse_quarter(self, bpage, scripcode):
        url = ("https://api.bseindia.com/BseIndiaAPI/api/"
               "CorporatesSHPSecuritybeta/w?scripcode=" + scripcode + "&qtrid=")
        res = self._bse_fetch(bpage, url)
        if not res.get("ok"):
            return None
        tbl = res["json"].get("Table") or []
        if not tbl:
            return None
        row = tbl[0]
        qid = row.get("Qtr_Id")
        qname = row.get("Fld_qtrname") or ""
        if qid is None:
            return None
        return {"qid": qid, "qname": qname,
                "as_of": _quarter_end_date(qname)}

    # ----- NSE ------------------------------------------------------------ #
    def _nse_ticker(self, npage, name, bse_name, fallback):
        query = re.sub(r"\s+(limited|ltd)\.?$", "", name, flags=re.IGNORECASE)
        url = ("https://www.nseindia.com/api/NextApi/globalSearch/equity?symbol="
               + urllib.parse.quote(query))
        targets = {_norm(name), _norm(bse_name)}
        for attempt in range(3):
            try:
                res = npage.evaluate(self._FETCH_JS, url)
            except Exception:
                # NSE occasionally reloads for its bot cookie; re-settle and retry
                try:
                    npage.goto("https://www.nseindia.com/",
                               wait_until="domcontentloaded")
                    npage.wait_for_timeout(1500)
                except Exception:
                    pass
                continue
            if res.get("ok"):
                data = (res["json"] or {}).get("data") or []
                eq = [d for d in data
                      if (d.get("series") or "").upper() == "EQ"] or data
                for d in eq:
                    if _norm(d.get("companyName", "")) in targets:
                        return (d.get("symbol") or "").strip()
                if eq:
                    return (eq[0].get("symbol") or "").strip()
                return fallback
            npage.wait_for_timeout(800)
        return fallback

    # ----- statement links + detailed shareholding ------------------------ #
    @staticmethod
    def _stmt_urls(scripcode, qid, qname):
        """Human-viewable statement pages (used as clickable links in Sheet 1)."""
        q = f"{float(qid):.2f}"
        enc = urllib.parse.quote(qname)
        base = "https://www.bseindia.com/corporates/"
        return (
            f"{base}ShpPromoterNGroup?scripcd={scripcode}&qtrid={q}&QtrName={enc}",
            f"{base}shpPublicShareholder?scripcd={scripcode}&qtrid={q}&QtrName={enc}",
        )

    @staticmethod
    def _largest_table(j):
        big, blen = None, -1
        for k, v in (j or {}).items():
            if isinstance(v, list) and len(v) > blen:
                big, blen = k, len(v)
        return (j.get(big) if big else []) or []

    @staticmethod
    def _share_table(j):
        """Prefer the SHP table that actually names holders and share counts."""
        best, n = None, -1
        for v in (j or {}).values():
            if (isinstance(v, list) and v and isinstance(v[0], dict)
                    and "Fld_ShareHolderName" in v[0]):
                if len(v) > n:
                    best, n = v, len(v)
        return best if best is not None else Pipeline._largest_table(j)

    @staticmethod
    def _shp_holder(x, category):
        nm = re.sub(r"\s+", " ", (x.get("Fld_ShareHolderName") or "").strip())
        pct = _to_float(x.get("Fld_TotalPercentageOf_A_B_C2"))
        shares = _to_float(x.get("Fld_TotalNoOfShares"))
        if not nm or pct is None or pct <= 0:
            return None
        return {
            "name": nm,
            "category": category,
            "pct": pct,
            "shares": shares,
            "holding_cr": None,
            "bought": 0,
            "sold": 0,
            "adj_shares": shares,
            "adj_pct": pct,
            "adj_holding_cr": None,
        }

    def _promoter_rows(self, bpage, scripcode, qid):
        """Named promoter / promoter-group holders with % and share count."""
        url = ("https://api.bseindia.com/BseIndiaAPI/api/"
               f"Corp_shpPromoterNGroup_ng/w?SCRIPCODE={scripcode}"
               f"&QtrCode={float(qid):.2f}")
        res = self._bse_fetch(bpage, url)
        out = []
        if res.get("ok"):
            for x in self._share_table(res["json"]):
                ty = (x.get("FLd_ShareholderType") or "").strip()
                if ty not in ("Promoter", "Promoter Group"):
                    continue
                row = self._shp_holder(x, ty)
                if row:
                    out.append(row)
        return out

    def _public_rows(self, bpage, scripcode, qid):
        """Named public holders with % and share count."""
        url = ("https://api.bseindia.com/BseIndiaAPI/api/"
               f"Corp_shpSec_SHPPubShold_ng/w?SCRIPCODE={scripcode}"
               f"&QtrCode={float(qid):.2f}")
        res = self._bse_fetch(bpage, url)
        out = []
        if res.get("ok"):
            for x in self._share_table(res["json"]):
                heading = (x.get("Fld_Level") or x.get("Fld_SubCategory") or "").strip()
                heading = re.sub(r"/+\s*$", "", heading).strip()
                row = self._shp_holder(x, heading)
                if row:
                    out.append(row)
        return out

    def _with_post_shp(self, bpage, npage, nse, scripcode, qtr, promoters,
                       public, mcap):
        holders = list(promoters) + list(public)
        _, extras, meta = fetch_and_adjust(
            bpage, npage, nse, scripcode, qtr.get("qname"), holders, mcap)
        if extras:
            public = list(public) + extras
        return promoters, public, meta

    # ----- orchestration -------------------------------------------------- #
    def _lookup_one_isolated(self, name, skip_tape=True):
        """Launch Chrome for one name, then close it (512 MB path)."""
        try:
            with sync_playwright() as p:
                browser, ctx, _label = _open_session(p, probe=False)
                try:
                    bpage = ctx.new_page()
                    try:
                        bpage.goto("https://www.bseindia.com/",
                                   wait_until="domcontentloaded",
                                   timeout=60000)
                        bpage.wait_for_timeout(800)
                    except Exception as e:
                        return {"found": False, "query": name,
                                "error": "Could not open BSE: " + _short_error(e)}
                    if _bse_blocked(bpage):
                        return {"found": False, "query": name,
                                "error": "BSE blocked the browser"}
                    return self._lookup_on_pages(
                        bpage, None, name, skip_tape=skip_tape)
                finally:
                    try:
                        ctx.close()
                    except Exception:
                        pass
                    try:
                        browser.close()
                    except Exception:
                        pass
        except Exception as e:
            return {"found": False, "query": name, "error": _short_error(e)}

    def _run_low_mem(self, dfrom, dto, out_path, threshold):
        self.log("Low-memory mode: IPO list over HTTP, Chrome opened and "
                 "closed per company, NSE block/bulk skipped.", stage="list")
        companies = self._listed_companies(None, dfrom, dto)
        self.log(f"Found {len(companies)} listed IPOs in range.", stage="list")
        qualified = []
        total = len(companies)
        fname = os.path.basename(out_path)
        cut = int(threshold) if threshold == int(threshold) else threshold
        title = f"IPO Companies >{cut}cr"
        for i, (name, _dt) in enumerate(companies, 1):
            self._check_cancel()
            self.log(f"[{i}/{total}] {name}", current=i, total=total,
                     stage="bse")
            rec = self._lookup_one_isolated(name, skip_tape=True)
            err = rec.get("error")
            if err and not rec.get("found"):
                self.log(f"    {err}", stage="bse")
                low = err.lower()
                if i == 1 and any(tok in low for tok in (
                        "executable", "could not start", "browser has been closed",
                        "chromium.launch", "host system is missing")):
                    raise RuntimeError(
                        "Chrome cannot run on this 512 MB instance. "
                        "Use this app on your PC (python app.py) or a 2 GB plan."
                    )
            if not rec.get("found"):
                self.log("    not found on BSE, skipping", stage="bse")
                continue
            mcap = rec.get("mcap")
            if mcap is None or mcap <= threshold:
                self.log(f"    Mcap {mcap} <= {threshold:.0f}cr, skipping",
                         stage="bse")
                continue
            if not rec.get("quarter"):
                self.log("    no shareholding quarter, skipping", stage="bse")
                continue
            rec["name"] = name
            qualified.append(rec)
            _write_workbook(qualified, out_path, sheet1_title=title)
            self.log(
                f"    QUALIFIED  Mcap {mcap:,.2f}cr  {rec.get('bse_ticker')}  "
                f"(promoters={len(rec.get('promoters') or [])}, "
                f"public={len(rec.get('public') or [])})",
                stage="bse", file=fname, count=len(qualified),
            )
        if not qualified:
            _write_workbook([], out_path, sheet1_title=title)
        self.log(f"Done. {len(qualified)} companies written.", stage="done",
                 file=fname, count=len(qualified))
        return qualified

    def run(self, dfrom, dto, out_path, mcap_min=None):
        threshold = MCAP_THRESHOLD if mcap_min is None else float(mcap_min)
        if _low_mem():
            return self._run_low_mem(dfrom, dto, out_path, threshold)
        qualified = []
        with sync_playwright() as p:
            self._check_cancel()
            browser, ctx, label = _open_session(p)
            self.log(f"Browser: {label}", stage="list")
            try:
                list_page = ctx.new_page()
                self.log("Fetching mainboard IPOs by listing date "
                         "(Chittorgarh)...", stage="list")
                companies = self._listed_companies(list_page, dfrom, dto)
                self.log(f"Found {len(companies)} listed IPOs in range.",
                         stage="list")
                list_page.close()
                self._check_cancel()

                self.log("Opening BSE for company lookup…", stage="bse")
                bpage = ctx.new_page()
                try:
                    bpage.goto("https://www.bseindia.com/",
                               wait_until="domcontentloaded", timeout=60000)
                    bpage.wait_for_timeout(1500)
                except Exception as e:
                    raise RuntimeError(
                        "Could not open BSE (common on Render datacenter IPs "
                        f"or a small instance). {e}"
                    ) from e
                if _bse_blocked(bpage):
                    raise RuntimeError(
                        "BSE blocked the browser after launch. "
                        "On Render, use the Docker runtime and a 2 GB instance."
                    )

                npage = None

                def ensure_nse():
                    nonlocal npage
                    if npage is not None:
                        return npage
                    npage = ctx.new_page()
                    try:
                        npage.goto("https://www.nseindia.com/",
                                   wait_until="domcontentloaded",
                                   timeout=45000)
                        npage.wait_for_timeout(1200)
                    except Exception:
                        pass
                    return npage

                self.log(f"Keeping companies with mcap > {threshold:.0f} cr.",
                         stage="bse")
                total = len(companies)
                for i, (name, dt) in enumerate(companies, 1):
                    self._check_cancel()
                    self.log(f"[{i}/{total}] {name}", current=i, total=total,
                             stage="bse")
                    info = self._bse_lookup(bpage, name)
                    if not info or not info["scripcode"]:
                        self.log(f"    not found on BSE, skipping", stage="bse")
                        continue
                    mcap = self._bse_mcap(bpage, info["scripcode"])
                    if mcap is None or mcap <= threshold:
                        self.log(f"    Mcap {mcap} <= {threshold:.0f}cr, skipping",
                                 stage="bse")
                        continue
                    qtr = self._bse_quarter(bpage, info["scripcode"])
                    if not qtr:
                        self.log("    no shareholding quarter, skipping", stage="bse")
                        continue
                    nse = self._nse_ticker(ensure_nse(), name, info["bse_name"],
                                           info["ticker"])
                    prom_url, pub_url = self._stmt_urls(
                        info["scripcode"], qtr["qid"], qtr["qname"])
                    promoters = _with_holdings(
                        self._promoter_rows(bpage, info["scripcode"], qtr["qid"]),
                        mcap)
                    public = _with_holdings(
                        self._public_rows(bpage, info["scripcode"], qtr["qid"]),
                        mcap)
                    promoters, public, meta = self._with_post_shp(
                        bpage, ensure_nse(), nse, info["scripcode"], qtr,
                        promoters, public, mcap)
                    qualified.append({
                        "name": name,
                        "bse_ticker": info["ticker"],
                        "scripcode": info["scripcode"],
                        "nse_ticker": nse,
                        "isin": info["isin"],
                        "mcap": mcap,
                        "quarter": _display_quarter(qtr["qname"]),
                        "quarter_end": _iso(meta.get("quarter_end") or qtr.get("as_of")),
                        "adjust_from": _iso(meta.get("adjust_from")),
                        "adjust_to": _iso(meta.get("adjust_to")),
                        "shares_outstanding": meta.get("shares_outstanding"),
                        "prom_url": prom_url,
                        "pub_url": pub_url,
                        "promoters": promoters,
                        "public": public,
                    })
                    self.log(f"    QUALIFIED  Mcap {mcap:,.2f}cr  {info['ticker']}  "
                             f"(promoters={len(promoters)}, public={len(public)})",
                             stage="bse")
                    bpage.wait_for_timeout(150)

                bpage.close()
                if npage is not None:
                    npage.close()
            finally:
                try:
                    ctx.close()
                except Exception:
                    pass
                try:
                    browser.close()
                except Exception:
                    pass

        self.log("Writing Excel workbook...", stage="excel")
        cut = int(threshold) if threshold == int(threshold) else threshold
        _write_workbook(qualified, out_path,
                        sheet1_title=f"IPO Companies >{cut}cr")
        self.log(f"Done. {len(qualified)} companies written.", stage="done")
        return qualified

    # ----- interactive search (one or many names, semicolon-separated) --- #
    def _lookup_on_pages(self, bpage, npage, name, skip_tape=False):
        """Resolve one name against pages already opened on BSE / NSE."""
        info = self._bse_lookup(bpage, name)
        if not info or not info["scripcode"]:
            return {"found": False, "query": name}

        mcap = self._bse_mcap(bpage, info["scripcode"])
        qtr = self._bse_quarter(bpage, info["scripcode"])
        nse = info["ticker"]
        if npage is not None and not skip_tape:
            try:
                nse = self._nse_ticker(npage, info["bse_name"],
                                       info["bse_name"], info["ticker"])
            except Exception:
                nse = info["ticker"]

        promoters, public, prom_url, pub_url, quarter = [], [], "", "", ""
        meta = {}
        if qtr:
            quarter = _display_quarter(qtr["qname"])
            prom_url, pub_url = self._stmt_urls(
                info["scripcode"], qtr["qid"], qtr["qname"])
            promoters = _with_holdings(
                self._promoter_rows(bpage, info["scripcode"], qtr["qid"]),
                mcap)
            public = _with_holdings(
                self._public_rows(bpage, info["scripcode"], qtr["qid"]),
                mcap)
            if skip_tape or npage is None:
                meta = {"quarter_end": qtr.get("as_of")}
            else:
                promoters, public, meta = self._with_post_shp(
                    bpage, npage, nse, info["scripcode"], qtr,
                    promoters, public, mcap)

        return {
            "found": True,
            "query": name,
            "name": info["bse_name"],
            "bse_ticker": info["ticker"],
            "scripcode": info["scripcode"],
            "nse_ticker": nse,
            "isin": info["isin"],
            "mcap": mcap,
            "quarter": quarter,
            "quarter_end": _iso(meta.get("quarter_end") or (qtr.get("as_of") if qtr else None)),
            "adjust_from": _iso(meta.get("adjust_from")),
            "adjust_to": _iso(meta.get("adjust_to")),
            "shares_outstanding": meta.get("shares_outstanding"),
            "prom_url": prom_url,
            "pub_url": pub_url,
            "promoters": promoters,
            "public": public,
        }

    def lookup_many(self, names):
        """Look up one or more companies in a single browser session.

        `names` is a string (`A; B; C`) or a list of terms. No market-cap
        filter — same as the interactive search.
        """
        terms = names if isinstance(names, (list, tuple)) else split_terms(names)
        if not terms:
            return []
        if _low_mem():
            out = []
            total = len(terms)
            for i, name in enumerate(terms, 1):
                self.log(f"[{i}/{total}] {name}", current=i, total=total,
                         stage="bse")
                out.append(self._lookup_one_isolated(name, skip_tape=True))
            return out
        with sync_playwright() as p:
            browser, ctx, _label = _open_session(p)
            try:
                bpage = ctx.new_page()
                bpage.goto("https://www.bseindia.com/",
                           wait_until="domcontentloaded")
                bpage.wait_for_timeout(1200)
                if _bse_blocked(bpage):
                    raise RuntimeError(
                        "BSE blocked the browser after launch. "
                        "Close other Chrome windows and try again."
                    )
                npage = ctx.new_page()
                try:
                    npage.goto("https://www.nseindia.com/",
                               wait_until="domcontentloaded")
                    npage.wait_for_timeout(1200)
                except Exception:
                    pass
                out = []
                total = len(terms)
                for i, name in enumerate(terms, 1):
                    self.log(f"[{i}/{total}] {name}", current=i, total=total,
                             stage="bse")
                    out.append(self._lookup_on_pages(bpage, npage, name))
                    bpage.wait_for_timeout(150)
                npage.close()
                bpage.close()
                return out
            finally:
                ctx.close()
                browser.close()

    def lookup_one(self, name):
        """Look up one company by name and return its full shareholding detail."""
        rows = self.lookup_many([name] if name else [])
        return rows[0] if rows else {"found": False, "query": name}


def _holder_view(row, mcap=None):
    """Normalise a holder row (dict or old list) for Excel / JSON."""
    if isinstance(row, dict):
        h = dict(row)
    else:
        h = {
            "name": row[0] if row else "",
            "category": row[1] if row and len(row) > 1 else "",
            "pct": row[2] if row and len(row) > 2 else None,
            "holding_cr": row[3] if row and len(row) > 3 else None,
            "shares": row[4] if row and len(row) > 4 else None,
        }
    if h.get("holding_cr") is None:
        h["holding_cr"] = _holding_cr(h.get("pct"), mcap)
    if h.get("adj_holding_cr") is None:
        h["adj_holding_cr"] = _holding_cr(h.get("adj_pct"), mcap)
    for k in ("bought", "sold"):
        h.setdefault(k, 0)
    return h


# --------------------------------------------------------------------------- #
# Excel writer
# --------------------------------------------------------------------------- #
def _write_workbook(companies, out_path, sheet1_title="IPO Companies >3000cr"):
    wb = Workbook()

    # ---- Sheet 1 ---- #
    ws = wb.active
    ws.title = sheet1_title[:31] or "Companies"
    headers = ["S.No", "Company Name", "BSE Ticker", "BSE Scrip Code",
               "NSE Ticker", "ISIN", "Live Mcap (Cr.)",
               "Latest Shareholding Quarter", "SHP as-of",
               "Shares Outstanding",
               "Adjusted from", "Adjusted to",
               "Promoter & Promoter Group Shareholding (link)",
               "Public Shareholder Shareholding (link)"]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    link_font = Font(color="0563C1", underline="single")

    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for i, comp in enumerate(companies, 1):
        r = i + 1
        vals = [i, comp["name"], comp["bse_ticker"], comp["scripcode"],
                comp["nse_ticker"], comp["isin"], comp["mcap"], comp["quarter"],
                comp.get("quarter_end") or "",
                comp.get("shares_outstanding"),
                comp.get("adjust_from") or "",
                comp.get("adjust_to") or ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = left if c == 2 else center
            if c == 7 and isinstance(v, (int, float)):
                cell.number_format = "#,##0.00"
            if c == 10 and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
        pc = ws.cell(row=r, column=13, value="Promoter & Promoter Group Statement")
        if comp.get("prom_url"):
            pc.hyperlink = comp["prom_url"]
        pc.font = link_font
        pc.alignment = left
        pc.border = border
        uc = ws.cell(row=r, column=14, value="Public Shareholder Statement")
        if comp.get("pub_url"):
            uc.hyperlink = comp["pub_url"]
        uc.font = link_font
        uc.alignment = left
        uc.border = border

    widths = [6, 42, 14, 12, 14, 16, 16, 16, 14, 18, 14, 14, 40, 36]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    if companies:
        ws.auto_filter.ref = f"A1:N{len(companies) + 1}"

    # ---- Sheet 2 ---- #
    ws2 = wb.create_sheet("Detailed Shareholding")
    h2 = ["Company Name", "Statement", "Category / Name of Shareholder",
          "Promoter Type / Public Heading",
          "Shares (SHP)", "Shareholding % (A+B+C2)", "Holding (Rs Cr.)",
          "Bought after SHP", "Sold after SHP",
          "Adj. shares", "Adj. %", "Adj. holding (Rs Cr.)"]
    prom_fill = PatternFill("solid", fgColor="E2EFDA")
    pub_fill = PatternFill("solid", fgColor="FCE4D6")
    comp_font = Font(bold=True, size=11, color="1F4E78")
    for c, h in enumerate(h2, 1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    def wrow(r, fill, vals, aligns):
        for c, (v, al) in enumerate(zip(vals, aligns), 1):
            cell = ws2.cell(row=r, column=c, value=v)
            cell.alignment = al
            cell.fill = fill
            cell.border = border
            if c in (5, 8, 9, 10) and isinstance(v, (int, float)):
                cell.number_format = "#,##0"
            if c in (7, 12) and isinstance(v, (int, float)):
                cell.number_format = "#,##0.00"
            if c in (6, 11) and isinstance(v, (int, float)):
                cell.number_format = "0.00"

    r = 2
    for comp in companies:
        start = r
        prom = comp.get("promoters") or []
        pub = comp.get("public") or []
        def pack(statement, row, aligns):
            h = _holder_view(row, comp.get("mcap"))
            return (
                [comp["name"], statement, h["name"], h["category"],
                 h["shares"], h["pct"], h["holding_cr"],
                 h["bought"], h["sold"],
                 h["adj_shares"], h["adj_pct"], h["adj_holding_cr"]],
                aligns,
            )

        if prom:
            for row in prom:
                vals, al = pack("Promoter & Promoter Group", row,
                                [left, left, left, center, center, center,
                                 center, center, center, center, center, center])
                wrow(r, prom_fill, vals, al)
                r += 1
        else:
            wrow(r, prom_fill,
                 [comp["name"], "Promoter & Promoter Group",
                  "No promoter / promoter group with shareholding > 0%",
                  "-", "-", "-", "-", "-", "-", "-", "-", "-"],
                 [left, left, left, center, center, center,
                  center, center, center, center, center, center])
            r += 1
        if pub:
            for row in pub:
                vals, al = pack("Public Shareholder", row,
                                [left, left, left, left, center, center,
                                 center, center, center, center, center, center])
                wrow(r, pub_fill, vals, al)
                r += 1
        else:
            wrow(r, pub_fill,
                 [comp["name"], "Public Shareholder",
                  "No named (non-bold) public shareholder with shareholding > 0%",
                  "-", "-", "-", "-", "-", "-", "-", "-", "-"],
                 [left, left, left, center, center, center,
                  center, center, center, center, center, center])
            r += 1
        ws2.cell(row=start, column=1).font = comp_font

    for c, w in enumerate([40, 26, 52, 40, 16, 14, 16, 16, 16, 16, 12, 18], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "A2"
    if r > 2:
        ws2.auto_filter.ref = f"A1:L{r - 1}"

    wb.save(out_path)


def run_pipeline(from_str, to_str, out_path, progress_cb=None, mcap_min=None,
                 cancel_cb=None):
    """from_str / to_str are 'YYYY-MM-DD' (HTML date input format)."""
    dfrom = datetime.datetime.strptime(from_str, "%Y-%m-%d").date()
    dto = datetime.datetime.strptime(to_str, "%Y-%m-%d").date()
    return Pipeline(progress_cb, cancel_cb).run(
        dfrom, dto, out_path, mcap_min=mcap_min)


def lookup_company(name):
    """Interactive lookup. `name` may be several companies separated by ';'."""
    terms = split_terms(name)
    return Pipeline().lookup_many(terms)
